"""Write FeatureImp validation reports, formatted SHAP and prediction workbooks, and
binary ensemble metrics.
"""
from __future__ import annotations

import os
import traceback
from datetime import datetime
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd

from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sklearn.metrics import f1_score, roc_auc_score

from .io_utils import _sanitize_excel_sheet_name
from .io_utils import _atomic_excel_writer


# For binary classification label encodings used throughout featureimp.
POSITIVE_CLASS_INDEX: int = 1
NEGATIVE_CLASS_INDEX: int = 0

def _print_validation_summary(validation_results):
    """Print a summary of validation results to console."""
    print("\n" + "="*80)
    print("VALIDATION SUMMARY")
    print("="*80)
    
    passed = sum(1 for check in validation_results['checks'].values() if check['status'] == 'PASS')
    warned = sum(1 for check in validation_results['checks'].values() if check['status'] == 'WARN')
    failed = sum(1 for check in validation_results['checks'].values() if check['status'] == 'FAIL')
    errors = sum(1 for check in validation_results['checks'].values() if check['status'] == 'ERROR')
    skipped = sum(1 for check in validation_results['checks'].values() if check['status'] == 'SKIP')
    
    print(f"Total checks: {len(validation_results['checks'])}")
    print(f"  âœ“ Passed: {passed}")
    print(f"  âš  Warnings: {warned}")
    print(f"  âœ— Failed: {failed}")
    print(f"  âœ— Errors: {errors}")
    print(f"  - Skipped: {skipped}")
    
    if validation_results['has_errors']:
        print(f"\n[ERROR] {len(validation_results['errors'])} validation errors detected!")
        for i, error in enumerate(validation_results['errors'][:5], 1):  
            print(f"  {i}. {error}")
        if len(validation_results['errors']) > 5:
            print(f"  ... and {len(validation_results['errors']) - 5} more errors")
    
    if validation_results['has_warnings']:
        print(f"\n[WARNING] {len(validation_results['warnings'])} validation warnings:")
        for i, warning in enumerate(validation_results['warnings'][:5], 1):  
            print(f"  {i}. {warning}")
        if len(validation_results['warnings']) > 5:
            print(f"  ... and {len(validation_results['warnings']) - 5} more warnings")
    
    if not validation_results['has_errors'] and not validation_results['has_warnings']:
        print("\n[OK] All validation checks passed!")
    
    print("="*80)

def _save_validation_report(validation_results, paths):
    """Save detailed validation report to file."""
    report_path = os.path.join(paths['analysis_folder'], 'VALIDATION_REPORT.txt')
    
    try:
        with open(report_path, 'w', encoding='utf-8') as f:
            f.write("="*80 + "\n")
            f.write("VALIDATION REPORT\n")
            f.write("="*80 + "\n\n")
            f.write(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n")
            
            # Summary
            passed = sum(1 for check in validation_results['checks'].values() if check['status'] == 'PASS')
            warned = sum(1 for check in validation_results['checks'].values() if check['status'] == 'WARN')
            failed = sum(1 for check in validation_results['checks'].values() if check['status'] == 'FAIL')
            errors = sum(1 for check in validation_results['checks'].values() if check['status'] == 'ERROR')
            skipped = sum(1 for check in validation_results['checks'].values() if check['status'] == 'SKIP')
            
            f.write("="*80 + "\n")
            f.write("SUMMARY\n")
            f.write("="*80 + "\n")
            f.write(f"Total checks: {len(validation_results['checks'])}\n")
            f.write(f"  [PASS] Passed: {passed}\n")
            f.write(f"  [WARN] Warnings: {warned}\n")
            f.write(f"  [FAIL] Failed: {failed}\n")
            f.write(f"  [ERROR] Errors: {errors}\n")
            f.write(f"  [SKIP] Skipped: {skipped}\n\n")
            
            # Detailed results
            f.write("="*80 + "\n")
            f.write("DETAILED RESULTS\n")
            f.write("="*80 + "\n\n")
            
            for check_name, check_result in validation_results['checks'].items():
                status = check_result['status']
                message = check_result['message']
                f.write(f"{check_name}:\n")
                f.write(f"  Status: {status}\n")
                f.write(f"  Message: {message}\n")
                # Add any additional metrics
                for key, value in check_result.items():
                    if key not in ['status', 'message']:
                        f.write(f"  {key}: {value}\n")
                f.write("\n")
            
            # Errors
            if validation_results['errors']:
                f.write("="*80 + "\n")
                f.write("ERRORS\n")
                f.write("="*80 + "\n")
                for i, error in enumerate(validation_results['errors'], 1):
                    f.write(f"{i}. {error}\n")
                f.write("\n")
            
            # Warnings
            if validation_results['warnings']:
                f.write("="*80 + "\n")
                f.write("WARNINGS\n")
                f.write("="*80 + "\n")
                for i, warning in enumerate(validation_results['warnings'], 1):
                    f.write(f"{i}. {warning}\n")
                f.write("\n")
        
        print(f"Validation report saved to: {report_path}")
    except Exception as e:
        print(f"Error saving validation report: {e}")
        traceback.print_exc()



# FILE I/O AND FORMATTING



def _consolidated_header(column):
    """Parse known metric prefixes without splitting underscores in group names."""
    text = str(column)
    section = ""
    for suffix, label in (("_OUTOFSAMPLE", "OUT-OF-SAMPLE"), ("_INSAMPLE", "IN-SAMPLE"),
                          ("_POOLED", "POOLED"), ("_MALE", "MALE"), ("_FEMALE", "FEMALE")):
        if text.endswith(suffix):
            text, section = text[:-len(suffix)], label
            break
    if text in ("", " "):
        return "", ""
    if text in ("q_value", "eta_squared"):
        return "Univariate", text
    metrics = ("MeanSHAP_Diff_Abs_Norm", "MeanSHAP_Diff", "MeanAbsSHAP", "MeanSHAP", "StdSHAP",
               "Direction_Agreement", "Stability_Ratio", "Rank_Correlation")
    for metric in metrics:
        prefix = metric + "_"
        if text.startswith(prefix):
            group = text[len(prefix):]
            normalized = group.startswith("Normalized_") or group.endswith("_Normalized")
            if group.startswith("Normalized_"):
                group = group[len("Normalized_"):]
            elif group.endswith("_Normalized"):
                group = group[:-len("_Normalized")]
            title = f"{group} ({section})" if section else group
            return title, metric + ("_Normalized" if normalized else "")
    return section, text


def _graphpad_group_names(groups):
    """Preserve original labels when sanitization would merge distinct groups."""
    labels = [str(group) for group in groups]
    sanitized = ["".join(c if c.isalnum() or c in ("-", "_") else "_" for c in label) for label in labels]
    # Excel cell headers support punctuation, so use complete originals on collision.
    return dict(zip(labels, labels if len(set(sanitized)) != len(sanitized) else sanitized))

def _write_consolidated_summary_sheet_with_formatting(
    writer: pd.ExcelWriter,
    df: pd.DataFrame,
    output_path: str,
    sheet_name: str,
    is_combined_insample_outofsample: bool = False,
) -> None:
    """Write one consolidated-summary sheet with the standard SHAP formatting."""

    is_sex_filtered = any(x in str(output_path).upper() for x in ['_MALE', '_FEMALE'])
    is_male_vs_female = 'MALE_VS_FEMALE' in str(output_path).upper()
    sheet_name = _sanitize_excel_sheet_name(sheet_name)

    # 1. Write the data WITHOUT the header, starting on row 3 to leave space.
    df.to_excel(writer, sheet_name=sheet_name, index=False, header=False, startrow=2)

    workbook = writer.book
    worksheet = workbook[sheet_name]

    # 2. Manually Write the 2-Level Header
    header_map: Dict[Any, Tuple[str, str]] = {}
    first_outofsample_col_idx = None

    for col_idx, col in enumerate(df.columns, 1):
        header_map[col] = _consolidated_header(col)
        if first_outofsample_col_idx is None and (
            (is_combined_insample_outofsample and str(col).endswith("_OUTOFSAMPLE")) or
            (is_male_vs_female and str(col).endswith("_FEMALE"))):
            first_outofsample_col_idx = col_idx

    for col_idx, col_name in enumerate(df.columns, 1):
        top_header, bottom_header = header_map[col_name]
        worksheet.cell(row=1, column=col_idx).value = top_header
        worksheet.cell(row=2, column=col_idx).value = bottom_header

    # 3. Merge Top-Level Header Cells
    header_row = worksheet[1]
    start_col_idx = 1
    current_header_val = header_row[0].value
    for col_idx in range(2, worksheet.max_column + 2):
        next_header_val = header_row[col_idx - 1].value if col_idx <= worksheet.max_column else None
        if next_header_val != current_header_val and current_header_val != '':
            end_col_idx = col_idx - 1
            if end_col_idx > start_col_idx:
                worksheet.merge_cells(start_row=1, start_column=start_col_idx, end_row=1, end_column=end_col_idx)

            merged_cell = worksheet.cell(row=1, column=start_col_idx)
            if is_sex_filtered or is_male_vs_female:
                merged_cell.alignment = Alignment(horizontal='left', vertical='center')
            else:
                merged_cell.alignment = Alignment(horizontal='center', vertical='center')

            start_col_idx = col_idx
            current_header_val = next_header_val
        elif next_header_val != current_header_val:
            start_col_idx = col_idx
            current_header_val = next_header_val

    # 4. Define Styles
    bold_font = Font(bold=True)
    italic_font = Font(italic=True)
    bold_italic_font = Font(bold=True, italic=True)
    thick_black_side = Side(style='thick', color='000000')
    thin_grey_side = Side(style='thin', color='D3D3D3')
    default_thin_border = Border(left=thin_grey_side, right=thin_grey_side, top=thin_grey_side, bottom=thin_grey_side)

    row1_alignment = Alignment(horizontal='left', vertical='center') if (is_sex_filtered or is_male_vs_female) else Alignment(horizontal='center', vertical='center')
    row2_alignment = Alignment(horizontal='left', vertical='center')

    for row in worksheet.iter_rows(min_row=1, max_row=1, min_col=1, max_col=worksheet.max_column):
        for cell in row:
            cell.font = bold_font
            cell.alignment = row1_alignment
            cell.border = Border(top=thick_black_side, bottom=thick_black_side, left=thick_black_side, right=thick_black_side)

    for row in worksheet.iter_rows(min_row=2, max_row=2, min_col=1, max_col=worksheet.max_column):
        for cell in row:
            cell.font = bold_font
            cell.alignment = row2_alignment
            cell.border = Border(top=thick_black_side, bottom=thick_black_side, left=thick_black_side, right=thick_black_side)

    # Column widths and freeze panes
    for col_idx, col in enumerate(df.columns, 1):
        col_letter = get_column_letter(col_idx)
        worksheet.column_dimensions[col_letter].width = 50 if col == 'feature_name' else 18
    worksheet.freeze_panes = 'A3'

    # 5. Apply Data Cell Formatting (Bolding, Borders, and Color Scale)
    df_cols = list(df.columns)
    feature_name_idx = df_cols.index('feature_name') + 1 if 'feature_name' in df_cols else -1

    global_impact_insample_idx = None
    global_impact_outofsample_idx = None
    global_impact_male_idx = None
    global_impact_female_idx = None
    for i, col in enumerate(df_cols):
        if col in ('Global_Importance', 'Global_Importance_INSAMPLE'):
            global_impact_insample_idx = i + 1
        elif col == 'Global_Importance_OUTOFSAMPLE':
            global_impact_outofsample_idx = i + 1
        elif col == 'Global_Importance_MALE':
            global_impact_male_idx = i + 1
        elif col == 'Global_Importance_FEMALE':
            global_impact_female_idx = i + 1

    feature_std_indices = {i + 1 for i, col in enumerate(df_cols) if str(col).startswith('FeatureStd_')}

    try:
        mean_shap_cols = [col for col in df.columns if str(col).startswith('MeanSHAP_') and not str(col).startswith('MeanSHAP_Diff') and not ('_Normalized' in str(col) or str(col).endswith('_Normalized'))]
        if mean_shap_cols:
            vmax = df[mean_shap_cols].abs().max().max() if not df[mean_shap_cols].empty else 1.0
            mean_shap_color_scale = ColorScaleRule(
                start_type='num', start_value=-vmax, start_color='FF0000',
                mid_type='num', mid_value=0, mid_color='FFFFFF',
                end_type='num', end_value=vmax, end_color='00FF00'
            )
            for col_idx, col_name in enumerate(df.columns, 1):
                col_name = str(col_name)
                if col_name.startswith('MeanSHAP_') and not col_name.startswith('MeanSHAP_Diff') and not ('_Normalized' in col_name or col_name.endswith('_Normalized')):
                    col_letter = worksheet.cell(row=3, column=col_idx).column_letter
                    cell_range = f"{col_letter}3:{col_letter}{worksheet.max_row}"
                    worksheet.conditional_formatting.add(cell_range, mean_shap_color_scale)

        mean_shap_norm_cols = [col for col in df.columns if str(col).startswith('MeanSHAP_Normalized_') or (str(col).startswith('MeanSHAP_') and '_Normalized' in str(col) and not str(col).startswith('MeanSHAP_Diff'))]
        if mean_shap_norm_cols:
            vmin_norm = df[mean_shap_norm_cols].min().min() if not df[mean_shap_norm_cols].empty else -0.1
            vmax_norm = df[mean_shap_norm_cols].max().max() if not df[mean_shap_norm_cols].empty else 0.1
            abs_max_norm = max(abs(vmin_norm), abs(vmax_norm))
            fixed_scale_max = max(0.1, abs_max_norm * 2) if abs_max_norm > 0 else 0.1
            mean_shap_norm_color_scale = ColorScaleRule(
                start_type='num', start_value=-fixed_scale_max, start_color='FF0000',
                mid_type='num', mid_value=0, mid_color='FFFFFF',
                end_type='num', end_value=fixed_scale_max, end_color='00FF00'
            )
            for col_idx, col_name in enumerate(df.columns, 1):
                col_name = str(col_name)
                if col_name.startswith('MeanSHAP_Normalized_') or (col_name.startswith('MeanSHAP_') and '_Normalized' in col_name and not col_name.startswith('MeanSHAP_Diff')):
                    col_letter = worksheet.cell(row=3, column=col_idx).column_letter
                    cell_range = f"{col_letter}3:{col_letter}{worksheet.max_row}"
                    worksheet.conditional_formatting.add(cell_range, mean_shap_norm_color_scale)

        mean_shap_diff_cols = [col for col in df.columns if str(col).startswith('MeanSHAP_Diff_') and not str(col).startswith('MeanSHAP_Diff_Abs_Norm_')]
        if mean_shap_diff_cols:
            vmax_diff = df[mean_shap_diff_cols].max().max() if not df[mean_shap_diff_cols].empty else 1.0
            if vmax_diff > 0:
                diff_color_scale = ColorScaleRule(
                    start_type='num', start_value=0, start_color='FFFFFF',
                    end_type='num', end_value=vmax_diff, end_color='007BA7'
                )
                for col_idx, col_name in enumerate(df.columns, 1):
                    col_name = str(col_name)
                    if col_name.startswith('MeanSHAP_Diff_') and not col_name.startswith('MeanSHAP_Diff_Abs_Norm_'):
                        col_letter = worksheet.cell(row=3, column=col_idx).column_letter
                        cell_range = f"{col_letter}3:{col_letter}{worksheet.max_row}"
                        worksheet.conditional_formatting.add(cell_range, diff_color_scale)

        mean_shap_diff_norm_cols = [col for col in df.columns if str(col).startswith('MeanSHAP_Diff_Abs_Norm_')]
        if mean_shap_diff_norm_cols:
            vmax_norm = df[mean_shap_diff_norm_cols].max().max() if not df[mean_shap_diff_norm_cols].empty else 1.0
            if vmax_norm > 0:
                diff_norm_color_scale = ColorScaleRule(
                    start_type='num', start_value=0, start_color='FFFFFF',
                    end_type='num', end_value=vmax_norm, end_color='8A2BE2'
                )
                for col_idx, col_name in enumerate(df.columns, 1):
                    if str(col_name).startswith('MeanSHAP_Diff_Abs_Norm_'):
                        col_letter = worksheet.cell(row=3, column=col_idx).column_letter
                        cell_range = f"{col_letter}3:{col_letter}{worksheet.max_row}"
                        worksheet.conditional_formatting.add(cell_range, diff_norm_color_scale)

        normalized_cols = [col for col in df.columns if ('_Normalized_' in str(col) or str(col).endswith('_Normalized')) and not str(col).startswith('MeanSHAP_Normalized_') and not (str(col).startswith('MeanSHAP_') and '_Normalized' in str(col))]
        if normalized_cols:
            normalized_color_scale = ColorScaleRule(
                start_type='num', start_value=0, start_color='FFFFFF',
                end_type='num', end_value=1.0, end_color='007BA7'
            )
            for col_idx, col_name in enumerate(df.columns, 1):
                col_name = str(col_name)
                if (('_Normalized_' in col_name or col_name.endswith('_Normalized')) and not col_name.startswith('MeanSHAP_Normalized_') and not (col_name.startswith('MeanSHAP_') and '_Normalized' in col_name)):
                    col_letter = worksheet.cell(row=3, column=col_idx).column_letter
                    cell_range = f"{col_letter}3:{col_letter}{worksheet.max_row}"
                    worksheet.conditional_formatting.add(cell_range, normalized_color_scale)

        cumulative_cols = [col for col in df.columns if 'Cumulative' in str(col)]
        if cumulative_cols:
            cumulative_color_scale = ColorScaleRule(
                start_type='num', start_value=0, start_color='FFFFFF',
                end_type='num', end_value=1.0, end_color='007BA7'
            )
            for col_idx, col_name in enumerate(df.columns, 1):
                if 'Cumulative' in str(col_name):
                    col_letter = worksheet.cell(row=3, column=col_idx).column_letter
                    cell_range = f"{col_letter}3:{col_letter}{worksheet.max_row}"
                    worksheet.conditional_formatting.add(cell_range, cumulative_color_scale)

        mean_abs_shap_cols = [col for col in df.columns if str(col).startswith('MeanAbsSHAP_') and not str(col).startswith('MeanAbsSHAP_Normalized_')]
        if mean_abs_shap_cols:
            vmax_abs = df[mean_abs_shap_cols].max().max() if not df[mean_abs_shap_cols].empty else 1.0
            if vmax_abs > 0:
                mean_abs_color_scale = ColorScaleRule(
                    start_type='num', start_value=0, start_color='FFFFFF',
                    end_type='num', end_value=vmax_abs, end_color='FFD700'
                )
                for col_idx, col_name in enumerate(df.columns, 1):
                    col_name = str(col_name)
                    if col_name.startswith('MeanAbsSHAP_') and not col_name.startswith('MeanAbsSHAP_Normalized_'):
                        col_letter = worksheet.cell(row=3, column=col_idx).column_letter
                        cell_range = f"{col_letter}3:{col_letter}{worksheet.max_row}"
                        worksheet.conditional_formatting.add(cell_range, mean_abs_color_scale)
    except Exception:
        pass

    for row in worksheet.iter_rows(min_row=3, max_row=worksheet.max_row, min_col=1, max_col=worksheet.max_column):
        for cell in row:
            cell.border = default_thin_border
            col_name = str(df_cols[cell.column - 1])

            is_normalized = '_Normalized' in col_name or col_name.endswith('_Normalized')
            is_cumulative = 'Cumulative' in col_name
            is_mean_shap_normalized = col_name.startswith('MeanSHAP_Normalized_') or (col_name.startswith('MeanSHAP_') and '_Normalized' in col_name)

            if is_mean_shap_normalized:
                cell.font = bold_italic_font
            elif is_normalized or is_cumulative:
                cell.font = italic_font
            elif col_name.startswith('MeanSHAP_') or col_name.startswith('FeatureMean_'):
                cell.font = bold_font

            if (is_combined_insample_outofsample or is_male_vs_female) and first_outofsample_col_idx and cell.column == first_outofsample_col_idx:
                cell.border = Border(left=thick_black_side, right=cell.border.right, top=cell.border.top, bottom=cell.border.bottom)
            elif cell.column == feature_name_idx:
                cell.border = Border(left=cell.border.left, right=thick_black_side, top=cell.border.top, bottom=cell.border.bottom)
            elif global_impact_insample_idx and cell.column == global_impact_insample_idx:
                cell.border = Border(left=cell.border.left, right=thick_black_side, top=cell.border.top, bottom=cell.border.bottom)
            elif global_impact_outofsample_idx and cell.column == global_impact_outofsample_idx:
                left_border = thick_black_side if ((is_combined_insample_outofsample or is_male_vs_female) and first_outofsample_col_idx and cell.column == first_outofsample_col_idx) else cell.border.left
                cell.border = Border(left=left_border, right=thick_black_side, top=cell.border.top, bottom=cell.border.bottom)
            elif global_impact_male_idx and cell.column == global_impact_male_idx:
                cell.border = Border(left=cell.border.left, right=thick_black_side, top=cell.border.top, bottom=cell.border.bottom)
            elif global_impact_female_idx and cell.column == global_impact_female_idx:
                left_border = thick_black_side if (is_male_vs_female and first_outofsample_col_idx and cell.column == first_outofsample_col_idx) else cell.border.left
                cell.border = Border(left=left_border, right=thick_black_side, top=cell.border.top, bottom=cell.border.bottom)
            elif cell.column in feature_std_indices:
                cell.border = Border(left=cell.border.left, right=thick_black_side, top=cell.border.top, bottom=cell.border.bottom)

            if cell.column != feature_name_idx and isinstance(cell.value, (int, float, np.integer, np.floating)) and not isinstance(cell.value, bool):
                cell.number_format = '#,##0.0000'

    if worksheet.max_row > 2 and worksheet.max_column > 0:
        try:
            last_col_letter = get_column_letter(worksheet.max_column)
            worksheet.auto_filter.ref = f"A2:{last_col_letter}{worksheet.max_row}"
        except Exception:
            pass



def save_consolidated_summary_sheets_with_formatting(
    sheets: Dict[str, pd.DataFrame],
    output_path: str,
    is_combined_insample_outofsample: bool = False,
    sheet_order: Optional[Sequence[str]] = None,
) -> None:
    """Save multiple consolidated-summary sheets into one workbook."""

    if not isinstance(sheets, dict) or not sheets:
        raise ValueError("sheets must be a non-empty dict of sheet_name -> DataFrame")

    ordered_names: List[str] = []
    if sheet_order:
        for n in sheet_order:
            if n in sheets and n not in ordered_names:
                ordered_names.append(n)
    for n in sheets.keys():
        if n not in ordered_names:
            ordered_names.append(n)

    with _atomic_excel_writer(output_path) as writer:
        for name in ordered_names:
            df = sheets.get(name)
            if df is None:
                raise ValueError(f'Required consolidated sheet is None: {name}')
            _write_consolidated_summary_sheet_with_formatting(
                writer,
                df,
                output_path=str(output_path),
                sheet_name=str(name),
                is_combined_insample_outofsample=bool(is_combined_insample_outofsample),
            )


def save_consolidated_summary_with_formatting(df, output_path, is_combined_insample_outofsample=False):
    """Write one formatted consolidated SHAP-summary workbook."""

    save_consolidated_summary_sheets_with_formatting(
        {"SHAP_Summary": df},
        output_path,
        is_combined_insample_outofsample=is_combined_insample_outofsample,
        sheet_order=["SHAP_Summary"],
    )

def save_predictions_with_formatting(df, output_path, extra_sheets=None, comparison_df=None,
                                    primary_sheet_name='Predictions',
                                    comparison_sheet_name='PredictionsComparison'):
    """
    Save predictions to Excel with consistent formatting.

    The first sheet is intended to contain the canonical weighted predictions.
    Optionally, a second fully formatted sheet can be written containing the
    unweighted predictions for direct comparison.
    """

    def _apply_predictions_formatting(worksheet, df_cols):
        """Apply shared prediction-sheet headers, widths, separators, and numeric
        formats.
        """
        bold_font = Font(bold=True)
        thick_black_side = Side(style='thick', color='000000')
        thin_grey_side = Side(style='thin', color='D3D3D3')
        default_thin_border = Border(left=thin_grey_side, right=thin_grey_side, top=thin_grey_side, bottom=thin_grey_side)
        thick_header_border = Border(top=thick_black_side, bottom=thick_black_side, left=thick_black_side, right=thick_black_side)

        insample_end_idx = None
        outofsample_end_idx = None
        first_data_col_idx = None
        last_col_idx = len(df_cols)

        for i, col in enumerate(df_cols):
            if col not in ['individual', 'group']:
                first_data_col_idx = i + 1
                break

        for i, col in enumerate(df_cols):
            if col.endswith('_insample'):
                insample_end_idx = i + 1

        for i, col in enumerate(df_cols):
            if col.endswith('_outofsample'):
                outofsample_end_idx = i + 1

        left_alignment = Alignment(horizontal='left', vertical='center')
        for col_idx in range(1, worksheet.max_column + 1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.font = bold_font
            cell.alignment = left_alignment
            cell.border = thick_header_border

        if first_data_col_idx:
            for row_idx in range(1, worksheet.max_row + 1):
                cell = worksheet.cell(row=row_idx, column=first_data_col_idx)
                current_border = cell.border if cell.border else default_thin_border
                cell.border = Border(left=thick_black_side, right=current_border.right, top=current_border.top, bottom=current_border.bottom)

        if insample_end_idx:
            for row_idx in range(1, worksheet.max_row + 1):
                cell = worksheet.cell(row=row_idx, column=insample_end_idx)
                current_border = cell.border if cell.border else default_thin_border
                cell.border = Border(left=current_border.left, right=thick_black_side, top=current_border.top, bottom=current_border.bottom)

        if outofsample_end_idx:
            for row_idx in range(1, worksheet.max_row + 1):
                cell = worksheet.cell(row=row_idx, column=outofsample_end_idx)
                current_border = cell.border if cell.border else default_thin_border
                cell.border = Border(left=current_border.left, right=thick_black_side, top=current_border.top, bottom=current_border.bottom)

        for row_idx in range(1, worksheet.max_row + 1):
            cell = worksheet.cell(row=row_idx, column=last_col_idx)
            current_border = cell.border if cell.border else default_thin_border
            cell.border = Border(left=current_border.left, right=thick_black_side, top=current_border.top, bottom=current_border.bottom)

        for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, min_col=1, max_col=worksheet.max_column):
            for cell in row:
                if (first_data_col_idx and cell.column == first_data_col_idx) or \
                   (insample_end_idx and cell.column == insample_end_idx) or \
                   (outofsample_end_idx and cell.column == outofsample_end_idx) or \
                   (cell.column == last_col_idx):
                    continue
                cell.border = default_thin_border

    primary_sheet_name = str(primary_sheet_name)[:31]
    comparison_sheet_name = str(comparison_sheet_name)[:31]

    with _atomic_excel_writer(output_path) as writer:
        df.to_excel(writer, sheet_name=primary_sheet_name, index=False)
        if comparison_df is not None:
            comparison_df.to_excel(writer, sheet_name=comparison_sheet_name, index=False)

        if extra_sheets:
            for extra_name, extra_df in (extra_sheets or {}).items():
                if extra_df is None:
                    continue
                try:
                    safe_name = str(extra_name)[:31]
                except Exception:
                    safe_name = 'Extra'
                try:
                    extra_df.to_excel(writer, sheet_name=safe_name, index=False)
                except Exception:
                    pass

        workbook = writer.book
        _apply_predictions_formatting(workbook[primary_sheet_name], list(df.columns))
        if comparison_df is not None:
            _apply_predictions_formatting(workbook[comparison_sheet_name], list(comparison_df.columns))

    print(f"Predictions saved with formatting to {output_path}")

def create_formatted_predictions_for_graphpad(all_predictions_df, training_group_a, training_group_b, 
                                             analysis_folder, file_suffix="", perm_idx=0):
    """
    Create a GraphPad-friendly predictions spreadsheet.
    
    The output is designed for copy/paste into GraphPad Prism:
    - One column per group/condition.
    - Values stacked vertically (no blank rows).
    - For training-group columns, probabilities are expressed as ``P(class == training_group_a)``
      so that columns are comparable.
    
    Args:
        all_predictions_df: DataFrame produced by the pipeline containing per-individual probabilities.
        training_group_a: Training group A label (class 0).
        training_group_b: Training group B label (class 1).
        analysis_folder: Folder where the output file will be written.
        file_suffix: Optional suffix appended to the output filename.
        perm_idx: Permutation index used when locating permutation outputs.
    
    Returns:
        Path to the saved ``.xlsx`` file.
    """
    # Remove blank rows (where group is NaN)
    all_predictions_df = all_predictions_df[all_predictions_df['group'].notna()].copy()
    
    # Sanitize training group names for column names
    safe_training_group_a = "".join([c if c.isalnum() or c in ('-','_') else '_' for c in str(training_group_a)])
    safe_training_group_b = "".join([c if c.isalnum() or c in ('-','_') else '_' for c in str(training_group_b)])
    
    # Identify training groups
    training_groups = {training_group_a, training_group_b}
    
    # Get all unique groups to identify non-training groups
    all_groups = all_predictions_df['group'].dropna().unique()
    non_training_groups = sorted([g for g in all_groups if g not in training_groups])
    group_labels = _graphpad_group_names(non_training_groups)
    
    # Build formatted data dictionary (each key is a column name, value is array of probabilities)
    formatted_data = {}
    
    # Training Group A: in-sample and out-of-sample
    group_a_mask = all_predictions_df['group'] == training_group_a
    if group_a_mask.sum() > 0:
        if f'Prob_{training_group_a}_insample' in all_predictions_df.columns:
            formatted_data[f'{safe_training_group_a}_training_group_A_insample'] = all_predictions_df.loc[group_a_mask, f'Prob_{training_group_a}_insample'].values
        if f'Prob_{training_group_a}_outofsample' in all_predictions_df.columns:
            formatted_data[f'{safe_training_group_a}_training_group_A_outofsample'] = all_predictions_df.loc[group_a_mask, f'Prob_{training_group_a}_outofsample'].values
    
    # Express group B predictions as P(group A), matching the other export columns.
    group_b_mask = all_predictions_df['group'] == training_group_b
    if group_b_mask.sum() > 0:
        if f'Prob_{training_group_a}_insample' in all_predictions_df.columns:
            formatted_data[f'{safe_training_group_b}_training_group_B_insample'] = all_predictions_df.loc[group_b_mask, f'Prob_{training_group_a}_insample'].values
        if f'Prob_{training_group_a}_outofsample' in all_predictions_df.columns:
            formatted_data[f'{safe_training_group_b}_training_group_B_outofsample'] = all_predictions_df.loc[group_b_mask, f'Prob_{training_group_a}_outofsample'].values
    
    # Non-training groups: combined probability of being training_group_a
    for nt_group in non_training_groups:
        nt_mask = all_predictions_df['group'] == nt_group
        if nt_mask.sum() > 0 and f'Prob_{training_group_a}' in all_predictions_df.columns:
            # Sanitize group name for column name (replace spaces/special chars with underscore)
            safe_group_name = group_labels[str(nt_group)]
            formatted_data[f'{safe_group_name}_NonTraining'] = all_predictions_df.loc[nt_mask, f'Prob_{training_group_a}'].values
    
    # Pad all columns to same length (max length across all columns)
    if len(formatted_data) > 0:
        max_len = max(len(v) for v in formatted_data.values())
        for key in formatted_data:
            current_len = len(formatted_data[key])
            if current_len < max_len:
                formatted_data[key] = np.append(formatted_data[key], [np.nan] * (max_len - current_len))
        
        formatted_main_df = pd.DataFrame(formatted_data)
    else:
        print("Warning: No data found for formatted output. Creating empty DataFrame.")
        formatted_main_df = pd.DataFrame()
    
    # Try to load permutation data
    candidates = [os.path.join(analysis_folder, 'PERMUTED_RESULTS', f'AVERAGED_PREDICTIONS_PERMUTED_{perm_idx}.xlsx'),
                  os.path.join(analysis_folder, f'PERMUTED_{perm_idx}_RESULTS', f'AVERAGED_PREDICTIONS_PERMUTED_{perm_idx}.xlsx')]
    perm_file = next((p for p in candidates if os.path.isfile(p)), candidates[0])
    
    formatted_perm_df = None
    if os.path.exists(perm_file):
        try:
            perm_df = pd.read_excel(perm_file)
            
            # Remove blank rows
            perm_df = perm_df[perm_df['group'].notna()].copy()
            
            # Build same structure for permutation data
            perm_formatted_data = {}
            
            # Training Group A
            perm_group_a_mask = perm_df['group'] == training_group_a
            if perm_group_a_mask.sum() > 0:
                if f'Prob_{training_group_a}_insample' in perm_df.columns:
                    perm_formatted_data[f'{safe_training_group_a}_training_group_A_insample'] = perm_df.loc[perm_group_a_mask, f'Prob_{training_group_a}_insample'].values
                if f'Prob_{training_group_a}_outofsample' in perm_df.columns:
                    perm_formatted_data[f'{safe_training_group_a}_training_group_A_outofsample'] = perm_df.loc[perm_group_a_mask, f'Prob_{training_group_a}_outofsample'].values
            
            # Training Group B (use Prob_{training_group_a}, not Prob_{training_group_b})
            perm_group_b_mask = perm_df['group'] == training_group_b
            if perm_group_b_mask.sum() > 0:
                if f'Prob_{training_group_a}_insample' in perm_df.columns:
                    perm_formatted_data[f'{safe_training_group_b}_training_group_B_insample'] = perm_df.loc[perm_group_b_mask, f'Prob_{training_group_a}_insample'].values
                if f'Prob_{training_group_a}_outofsample' in perm_df.columns:
                    perm_formatted_data[f'{safe_training_group_b}_training_group_B_outofsample'] = perm_df.loc[perm_group_b_mask, f'Prob_{training_group_a}_outofsample'].values
            
            # Non-training groups
            perm_all_groups = perm_df['group'].dropna().unique()
            perm_non_training_groups = sorted([g for g in perm_all_groups if g not in training_groups])
            perm_group_labels = _graphpad_group_names(perm_non_training_groups)
            for nt_group in perm_non_training_groups:
                perm_nt_mask = perm_df['group'] == nt_group
                if perm_nt_mask.sum() > 0 and f'Prob_{training_group_a}' in perm_df.columns:
                    safe_group_name = perm_group_labels[str(nt_group)]
                    perm_formatted_data[f'{safe_group_name}_NonTraining'] = perm_df.loc[perm_nt_mask, f'Prob_{training_group_a}'].values
            
            # Pad permutation columns to same length
            if len(perm_formatted_data) > 0:
                perm_max_len = max(len(v) for v in perm_formatted_data.values())
                for key in perm_formatted_data:
                    current_len = len(perm_formatted_data[key])
                    if current_len < perm_max_len:
                        perm_formatted_data[key] = np.append(perm_formatted_data[key], [np.nan] * (perm_max_len - current_len))
                
                # Add "Perm_" prefix to all permutation column names
                perm_formatted_data_prefixed = {f'Perm_{key}': value for key, value in perm_formatted_data.items()}
                formatted_perm_df = pd.DataFrame(perm_formatted_data_prefixed)
            else:
                print("Warning: No permutation data found for formatted output.")
                
        except Exception as e:
            print(f"Warning: Could not load permutation file for formatted output: {e}")
            traceback.print_exc()
    
    # Combine main and permutation blocks
    if formatted_perm_df is not None and len(formatted_perm_df) > 0:
        # Ensure same number of rows (pad if needed)
        main_len = len(formatted_main_df)
        perm_len = len(formatted_perm_df)
        
        if main_len != perm_len:
            if main_len < perm_len:
                # Pad main with NaN rows
                pad_rows = pd.DataFrame(np.nan, index=range(perm_len - main_len), columns=formatted_main_df.columns)
                formatted_main_df = pd.concat([formatted_main_df, pad_rows], ignore_index=True)
            else:
                # Pad perm with NaN rows
                pad_rows = pd.DataFrame(np.nan, index=range(main_len - perm_len), columns=formatted_perm_df.columns)
                formatted_perm_df = pd.concat([formatted_perm_df, pad_rows], ignore_index=True)
        
        # Add separator column
        formatted_main_df['___SEPARATOR___'] = np.nan
        # Combine horizontally
        formatted_df = pd.concat([formatted_main_df, formatted_perm_df], axis=1)
    else:
        formatted_df = formatted_main_df
    
    # Save formatted output
    formatted_path = os.path.join(analysis_folder, f'AVERAGED_PREDICTIONS_FORMATTED_FOR_GRAPHPAD{file_suffix}.xlsx')
    with _atomic_excel_writer(formatted_path) as writer:
        formatted_df.to_excel(writer, index=False, sheet_name='Sheet1')
        workbook = writer.book
        worksheet = workbook['Sheet1']
        
        # Left-justify headers (row 1)
        for cell in worksheet[1]:
            cell.alignment = Alignment(horizontal='left', vertical='center')
    
    print(f"\nFormatted predictions for GraphPad saved to {formatted_path}")
    if formatted_perm_df is not None and len(formatted_perm_df) > 0:
        print(f"  - Includes main analysis and permutation {perm_idx} data")
    else:
        print(f"  - Includes main analysis only (permutation file not found)")
    
    return formatted_path

def calculate_ensemble_metrics(y_true, insample_probas, oof_probas, num_classes, le):
    """
    Compute binary classification metrics for in-sample and out-of-sample predictions.
    
    Args:
        y_true: True encoded labels (0/1) for training-group samples.
        insample_probas: In-sample predicted probabilities, shape ``(n_samples, 2)``.
        oof_probas: Out-of-fold predicted probabilities, shape ``(n_samples, 2)``.
        num_classes: Number of classes (must be 2).
        le: Label encoder instance (not required for computation; kept for signature consistency).
    
    Returns:
        dict with keys: ``f1_insample``, ``f1_outofsample``, ``auc_insample``, ``auc_outofsample``.
    
    Raises:
        ValueError: If inputs are not binary or shapes do not align.
    
    Notes:
        This helper assumes the fixed mapping used throughout the script:
        class 0 -> training_group_a
        class 1 -> training_group_b
    """
    metrics = {}

    y_true = np.asarray(y_true, dtype=int)
    insample_probas = np.asarray(insample_probas)
    oof_probas = np.asarray(oof_probas)

    if num_classes != 2:
        raise ValueError(f"calculate_ensemble_metrics expects binary classification (num_classes=2). Got {num_classes}.")

    if insample_probas.ndim != 2 or oof_probas.ndim != 2:
        raise ValueError(
            f"Expected 2D probability arrays. Got insample_probas.ndim={insample_probas.ndim}, "
            f"oof_probas.ndim={oof_probas.ndim}."
        )

    if insample_probas.shape[1] != num_classes or oof_probas.shape[1] != num_classes:
        raise ValueError(
            f"Probability arrays must have shape (n_samples, {num_classes}). "
            f"Got insample_probas.shape={insample_probas.shape}, oof_probas.shape={oof_probas.shape}."
        )

    if insample_probas.shape[0] != y_true.shape[0] or oof_probas.shape[0] != y_true.shape[0]:
        raise ValueError(
            f"y_true and probability arrays must align on n_samples. "
            f"Got y_true={y_true.shape[0]}, insample={insample_probas.shape[0]}, oof={oof_probas.shape[0]}."
        )

    unique_classes = np.unique(y_true)
    if unique_classes.size != 2:
        raise ValueError(
            f"Cannot compute binary metrics: y_true contains {unique_classes.size} unique classes: {unique_classes}"
        )

    # Convert probabilities to class predictions for F1 score
    insample_pred_classes = np.argmax(insample_probas, axis=1)
    oof_pred_classes = np.argmax(oof_probas, axis=1)

    # Calculate F1 scores (binary classification) â€” explicitly define the positive label
    metrics['f1_insample'] = f1_score(y_true, insample_pred_classes, average='binary', pos_label=POSITIVE_CLASS_INDEX)
    metrics['f1_outofsample'] = f1_score(y_true, oof_pred_classes, average='binary', pos_label=POSITIVE_CLASS_INDEX)

    # AUC for binary: use positive class probabilities
    metrics['auc_insample'] = roc_auc_score(y_true, insample_probas[:, POSITIVE_CLASS_INDEX])
    metrics['auc_outofsample'] = roc_auc_score(y_true, oof_probas[:, POSITIVE_CLASS_INDEX])

    return metrics
