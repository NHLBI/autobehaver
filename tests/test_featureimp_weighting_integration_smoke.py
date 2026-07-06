import os
import sys
import types
from pathlib import Path

import numpy as np
import pandas as pd

from openpyxl import load_workbook


def _touch(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("dummy", encoding="utf-8")


def _write_memmap(path: Path, arr: np.ndarray) -> str:
    path.parent.mkdir(parents=True, exist_ok=True)
    arr = np.asarray(arr, dtype=np.float32)
    mm = np.memmap(str(path), dtype=np.float32, mode="w+", shape=arr.shape)
    mm[:] = arr[:]
    mm.flush()
    del mm
    return str(path)


def test_weighting_stack_smoke(tmp_path: Path, monkeypatch):
    """End-to-end smoke for Phase2->Phase3->Phase4->Finalize outputs.

    This test does NOT train real XGBoost models. It monkeypatches model loading
    and SHAP computation so it runs fast while still exercising:
      - subset weight IO
      - logit aggregation
      - weighted SHAP aggregation
      - consolidated workbook + prediction workbook schemas
    """

    # ------------------------------------------------------------------
    # 0) Tiny synthetic dataset (training groups only; no non-training plots)
    # ------------------------------------------------------------------
    rng = np.random.RandomState(0)
    feature_names = ["f0", "f1", "f2"]

    df = pd.DataFrame(
        {
            "individual": [f"i{i}" for i in range(6)],
            "group": ["A", "B", "A", "B", "A", "B"],
        }
    )

    X_train = pd.DataFrame(rng.randn(len(df), len(feature_names)), columns=feature_names)
    X_all = X_train.copy()

    # Binary labels (A->0, B->1)
    y = np.array([0, 1, 0, 1, 0, 1], dtype=int)

    # The pipeline stratifies folds using a combined sex/class label.
    # For a unit test, using y is sufficient as a 2-stratum label.
    strat = y.copy()

    n_splits = 2
    GLOBAL_RANDOM_SEED = 123
    h_hit = 1

    # ------------------------------------------------------------------
    # 1) Create subset performance CSVs + compute weights (Phase 2)
    # ------------------------------------------------------------------
    analysis_folder = tmp_path
    (analysis_folder / "COMMITTEE_SUBSETS").mkdir(parents=True, exist_ok=True)

    real_perf = pd.DataFrame(
        {
            "h_hit": [h_hit, h_hit],
            "subset_idx": [0, 1],
            "delta_norm": [0.50, 0.10],
            "subset_complete": [True, True],
        }
    )
    real_perf_path = analysis_folder / "COMMITTEE_SUBSETS" / "SUBSET_PERFORMANCE_REAL.csv"
    real_perf.to_csv(real_perf_path, index=False)

    # Perm pooled table (multiple perm_idx rows) to set tau
    perm_perf = pd.DataFrame(
        {
            "perm_idx": [0, 1, 0, 1],
            "h_hit": [h_hit, h_hit, h_hit, h_hit],
            "subset_idx": [0, 0, 1, 1],
            "delta_norm": [0.05, 0.02, 0.01, 0.04],
            "subset_complete": [True, True, True, True],
        }
    )

    perm_base = analysis_folder / "PERMUTED_RESULTS"
    (perm_base / "COMMITTEE_SUBSETS").mkdir(parents=True, exist_ok=True)
    perm_perf_path = perm_base / "COMMITTEE_SUBSETS" / "SUBSET_PERFORMANCE_PERMUTED_ALL.csv"
    perm_perf.to_csv(perm_perf_path, index=False)

    from autobehaver.postinf.featureimp.subset_weighting import (
        SubsetWeightingConfig,
        compute_and_write_subset_weights,
    )

    cfg = SubsetWeightingConfig(tau_quantile=0.95, gamma=2.0)
    manifest = compute_and_write_subset_weights(
        analysis_folder=str(analysis_folder),
        permuted_base_folder=str(perm_base),
        cfg=cfg,
        verbose=False,
    )
    assert manifest is not None

    weights_csv = analysis_folder / "COMMITTEE_SUBSETS" / "SUBSET_WEIGHTS_REAL.csv"
    assert weights_csv.exists()

    from autobehaver.postinf.featureimp.logit_ensemble import load_subset_weights_mapping

    subset_weights = load_subset_weights_mapping(str(weights_csv))
    # Expect subset 0 to have larger weight than subset 1.
    assert subset_weights[(h_hit, 0)] > subset_weights[(h_hit, 1)]

    # ------------------------------------------------------------------
    # 2) Monkeypatch XGBoost in logit_ensemble so we can aggregate without real models
    # ------------------------------------------------------------------
    from autobehaver.postinf.featureimp import logit_ensemble as le

    class DummyModel:
        def __init__(self):
            self._margin = 0.0

        def load_model(self, path):
            # Encode deterministic margin from the filename.
            base = os.path.basename(str(path))
            margin = 0.0
            for part in base.replace(".", "_").split("_"):
                if part.startswith("s") and part[1:].isdigit():
                    margin += float(int(part[1:]))
                if part.startswith("f") and part[1:].isdigit():
                    margin += float(int(part[1:])) / 10.0
            self._margin = margin

        def predict(self, X, output_margin=False):
            return np.full((len(X),), self._margin, dtype=np.float64)

    orig_xgb_cls = le.xgb.XGBClassifier
    le.xgb.XGBClassifier = DummyModel

    try:
        # Build dummy fold models + performance records
        records = []
        model_dir = analysis_folder / "MODELS"
        subsets = {0: ["f0", "f1"], 1: ["f1", "f2"]}
        for subset_idx, feats in subsets.items():
            for fold_idx in range(n_splits):
                mp = model_dir / f"h{h_hit}_s{subset_idx}_f{fold_idx}.json"
                _touch(mp)
                records.append(
                    {
                        "h_hit": h_hit,
                        "subset_idx": subset_idx,
                        "fold_idx": fold_idx,
                        "model_path": str(mp),
                        "specialist_features": feats,
                        "is_permuted": False,
                    }
                )

        out = le.compute_logit_ensemble_predictions_from_saved_models(
            performance_records=records,
            X_train_full=X_train,
            X_all_groups=X_all,
            y_encoded_train=y,
            sex_group_encoded=strat,
            num_classes=2,
            n_splits=n_splits,
            GLOBAL_RANDOM_SEED=GLOBAL_RANDOM_SEED,
            subset_weights=subset_weights,
            verbose=False,
        )

        assert out.oof_probas_weighted is not None
        assert out.insample_probas_weighted is not None
        assert out.all_groups_probas_weighted is not None

        # Weighted should differ from unweighted in this test (weights are not equal).
        assert not np.allclose(out.oof_probas_weighted, out.oof_probas)

    finally:
        le.xgb.XGBClassifier = orig_xgb_cls

    # ------------------------------------------------------------------
    # 3) Weighted SHAP aggregation smoke (monkeypatch shap + xgboost)
    # ------------------------------------------------------------------
    from autobehaver.postinf.featureimp import shap_ensemble as se

    class DummyModel2:
        def __init__(self):
            self._value = 0.0

        def load_model(self, path):
            base = os.path.basename(str(path))
            val = 0.0
            for part in base.replace(".", "_").split("_"):
                if part.startswith("f") and part[1:].isdigit():
                    val += float(int(part[1:]))
            self._value = val

    dummy_shap = types.ModuleType("shap")

    class DummyTreeExplainer:
        def __init__(self, model):
            self.model = model

        def shap_values(self, X):
            # deterministic per-model SHAP values
            return np.ones((len(X), X.shape[1]), dtype=np.float32) * float(getattr(self.model, "_value", 0.0))

    dummy_shap.TreeExplainer = DummyTreeExplainer

    orig_shap_mod = sys.modules.get("shap")
    sys.modules["shap"] = dummy_shap

    orig_xgb_cls2 = se.xgb.XGBClassifier
    se.xgb.XGBClassifier = DummyModel2

    try:
        from autobehaver.postinf.featureimp.hhit import make_hhit_folds

        fold_seed = GLOBAL_RANDOM_SEED + h_hit * 1000
        folds = make_hhit_folds(X_train, strat, n_splits=n_splits, random_state=fold_seed)
        fold_indices_by_h_hit = {h_hit: folds}

        oof_w, ins_w, all_w, feat_wsum, diag = se.compute_weighted_shap_from_saved_models(
            performance_records=records,
            subset_weights=subset_weights,
            fold_indices_by_h_hit=fold_indices_by_h_hit,
            X_all_groups=X_all,
            all_feature_names=feature_names,
            n_train=len(X_train),
            n_splits=n_splits,
            verbose=False,
        )

        assert diag["n_subsets_used"] > 0
        assert all_w.shape == (len(X_all), len(feature_names))
        assert feat_wsum.shape == (len(feature_names),)

        # Also compute an unweighted-equitable view (all weights = 1)
        ones = {(h_hit, 0): 1.0, (h_hit, 1): 1.0}
        oof_u, ins_u, all_u, feat_usum, _ = se.compute_weighted_shap_from_saved_models(
            performance_records=records,
            subset_weights=ones,
            fold_indices_by_h_hit=fold_indices_by_h_hit,
            X_all_groups=X_all,
            all_feature_names=feature_names,
            n_train=len(X_train),
            n_splits=n_splits,
            verbose=False,
        )

        # Compute a simple ensemble-consistent view: numerator/total_weight_sum.
        # Since our equitable output is num/feat_sum, multiply by feat_sum/total.
        total_w = float(np.sum([subset_weights[(h_hit, 0)], subset_weights[(h_hit, 1)]]))
        p_feat = (feat_wsum / total_w).astype(np.float32)
        all_cons = all_w * p_feat[None, :]

    finally:
        se.xgb.XGBClassifier = orig_xgb_cls2
        if orig_shap_mod is not None:
            sys.modules["shap"] = orig_shap_mod
        else:
            del sys.modules["shap"]

    # ------------------------------------------------------------------
    # 4) Finalize: consolidated SHAP workbook should include the expected tabs
    # ------------------------------------------------------------------
    from autobehaver.postinf.featureimp.committee import FixedBinaryLabelEncoder
    from autobehaver.postinf.featureimp.finalize import _aggregate_and_summarize_results

    le_bin = FixedBinaryLabelEncoder("A", "B")

    # Main SHAP view = weighted equitable
    weighted_path = _write_memmap(analysis_folder / "all_groups_avg_shap_values.mmap", all_w)
    unweighted_path = _write_memmap(analysis_folder / "unweighted_view.mmap", all_u)
    consistent_path = _write_memmap(analysis_folder / "consistent_view.mmap", all_cons)

    # Global-importance overrides (simple, deterministic)
    gi_unweighted = np.mean(np.abs(all_u), axis=0)
    gi_weighted_raw = np.mean(np.abs(all_w), axis=0)
    gi_perm = np.full_like(gi_weighted_raw, 0.01)
    gi_excess = np.clip(gi_weighted_raw - gi_perm, 0, None)

    overrides = {
        "Global_Importance": gi_excess,
        "Global_Importance_Unweighted": gi_unweighted,
        "Global_Importance_WeightedRaw": gi_weighted_raw,
        "Global_Importance_PermBaseline": gi_perm,
    }

    paths = {"analysis_folder": str(analysis_folder), "file_suffix": ""}

    _aggregate_and_summarize_results(
        df=pd.concat([df, X_train], axis=1),
        selected_features=feature_names,
        performance_results=[],
        all_shap_values=[weighted_path],
        le=le_bin,
        paths=paths,
        global_importance_overrides=overrides,
        extra_consolidated_shap_views={
            "UnweightedEquitable": unweighted_path,
            "WeightedEnsembleConsistent": consistent_path,
        },
    )

    shap_xlsx = analysis_folder / "SHAP_CONSOLIDATED_SUMMARY.xlsx"
    assert shap_xlsx.exists()

    wb = load_workbook(shap_xlsx, read_only=True)
    sheets = set(wb.sheetnames)
    assert "Weighted SHAP" in sheets
    assert "Unweighted SHAP" in sheets
    assert "SHAP_Summary" not in sheets
    assert "WeightedEquitable" not in sheets
    assert "UnweightedEquitable" not in sheets
    assert "WeightedEnsembleConsistent" not in sheets

    wb_fmt = load_workbook(shap_xlsx, read_only=False, data_only=True)
    assert wb_fmt["Weighted SHAP"]["B3"].number_format == "#,##0.0000"

    ranking_csv = analysis_folder / "GLOBAL_IMPORTANCE_RANKING.csv"
    assert ranking_csv.exists()
    ranking = pd.read_csv(ranking_csv)
    assert "Global_Importance" in ranking.columns
    assert "Global_Importance_WeightedRaw" not in ranking.columns
    assert "Global_Importance_PermBaseline" not in ranking.columns
    assert (ranking["Global_Importance"] >= 0).all()

    # ------------------------------------------------------------------
    # 5) Predictions workbook schema smoke
    # ------------------------------------------------------------------
    from autobehaver.postinf.featureimp.reporting import save_predictions_with_formatting

    prob_cols = ["Prob_A", "Prob_B"]

    all_pred_weighted = pd.DataFrame(
        {
            "individual": df["individual"],
            "group": df["group"],
            "Prob_A_insample": out.insample_probas_weighted[:, 0],
            "Prob_B_insample": out.insample_probas_weighted[:, 1],
            "Prob_A_outofsample": out.oof_probas_weighted[:, 0],
            "Prob_B_outofsample": out.oof_probas_weighted[:, 1],
            "Prob_A": out.all_groups_probas_weighted[:, 0],
            "Prob_B": out.all_groups_probas_weighted[:, 1],
        }
    )

    all_pred_unweighted = pd.DataFrame(
        {
            "individual": df["individual"],
            "group": df["group"],
            "Prob_A_insample": out.insample_probas[:, 0],
            "Prob_B_insample": out.insample_probas[:, 1],
            "Prob_A_outofsample": out.oof_probas[:, 0],
            "Prob_B_outofsample": out.oof_probas[:, 1],
            "Prob_A": out.all_groups_probas[:, 0],
            "Prob_B": out.all_groups_probas[:, 1],
        }
    )

    preds_xlsx = analysis_folder / "AVERAGED_PREDICTIONS_FULL.xlsx"
    save_predictions_with_formatting(
        all_pred_weighted,
        str(preds_xlsx),
        comparison_df=all_pred_unweighted,
        primary_sheet_name="WeightedPredictions",
        comparison_sheet_name="UnweightedPredictions",
    )

    assert preds_xlsx.exists()
    wb2 = load_workbook(preds_xlsx, read_only=True)
    assert wb2.sheetnames[0] == "WeightedPredictions"
    assert "UnweightedPredictions" in wb2.sheetnames
