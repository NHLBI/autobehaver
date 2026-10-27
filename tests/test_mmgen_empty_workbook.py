"""Verify Excel reports remain valid when no summaries are eligible."""

import pandas as pd
import pytest
from openpyxl import load_workbook

from autobehaver.postinf.mmgen.feature_matrix import save_to_excel_sheets


@pytest.mark.parametrize('sheets, order', [
    ({}, None),
    ({'Absent': None, 'Empty': pd.DataFrame(columns=['value'])}, None),
    ({'Available': pd.DataFrame({'value': [1]})}, ['Missing']),
])
def test_unavailable_summaries_produce_a_readable_no_data_report(tmp_path, sheets, order):
    output_path = tmp_path / 'summary.xlsx'
    save_to_excel_sheets(sheets, output_path, sheet_order=order)

    workbook = load_workbook(output_path)
    try:
        assert workbook.sheetnames == ['No data']
        sheet = workbook['No data']
        assert sheet.sheet_state == 'visible'
        assert list(sheet.values) == [
            ('Status',),
            ('No eligible summary data were available for this report.',),
        ]
        assert sheet['A1'].alignment.horizontal == 'left'
    finally:
        workbook.close()


def test_populated_summary_preserves_order_names_and_values(tmp_path):
    output_path = tmp_path / 'summary.xlsx'
    sheets = {
        'Pose/Orient': pd.DataFrame({'individual': ['fly1', 'fly2'], 'value': [0, 2]}),
        'Time budget': pd.DataFrame({'individual': ['fly1'], 'value': [0.25]}),
        'Empty': pd.DataFrame(columns=['value']),
    }
    save_to_excel_sheets(sheets, output_path, sheet_order=['Time budget', 'Empty', 'Pose/Orient'])

    workbook = load_workbook(output_path)
    try:
        assert workbook.sheetnames == ['Time budget', 'Pose_Orient']
        assert list(workbook['Time budget'].values) == [('individual', 'value'), ('fly1', 0.25)]
        assert list(workbook['Pose_Orient'].values) == [('individual', 'value'), ('fly1', 0), ('fly2', 2)]
        assert all(cell.alignment.horizontal == 'left' for cell in workbook['Pose_Orient'][1])
    finally:
        workbook.close()
